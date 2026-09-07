"""Document creation and editing — with images as first-class content.

God-Agent could produce text and it could produce images, but it had no way to
put them together. "Screenshot the dashboard, search for the error, and write
me a report" ended with a pile of loose files.

A document here is an ordered list of **blocks** (heading / text / code /
quote / image / page_break) stored as JSON, alongside an assets directory
holding the images. Editing is by block index, so the agent can revise a
specific paragraph without regenerating the whole thing.

Render targets:

  md    pure stdlib — always available
  html  pure stdlib, single self-contained file (images base64-embedded)
  pdf   HTML rendered through the headless Chromium the browser tools
        already use — no extra dependency
  docx  python-docx (optional extra: pip install 'god-agent[docs]')

Because rendering is separated from the model, one document renders to every
format, and re-rendering after an edit is cheap.

Image embeds are copied into the document's assets directory, so a document
stays portable when you move it — the screenshots it cites travel with it.
"""
from __future__ import annotations

import base64
import html as _html
import json
import mimetypes
import os
import shutil
import time
from typing import Any, Optional

from ..runtime import get_runtime

SCHEMA_VERSION = 1
SUFFIX = ".gdoc.json"
BLOCK_KINDS = ("heading", "text", "code", "quote", "image", "page_break")
RENDER_FORMATS = ("md", "html", "pdf", "docx")


class DocError(RuntimeError):
    pass


# ---------------------------------------------------------------------------
# Paths and persistence
# ---------------------------------------------------------------------------
def _paths(path: str) -> tuple[str, str]:
    """Return (json_path, assets_dir) for a document name."""
    p = os.path.expanduser(str(path))
    for suf in (SUFFIX, ".gdoc"):
        if p.endswith(suf):
            base = p[: -len(suf)]
            break
    else:
        base = p
    return base + SUFFIX, base + ".assets"


def _load(json_path: str) -> dict:
    if not os.path.isfile(json_path):
        raise DocError(f"no such document: {json_path}")
    try:
        with open(json_path, "r", encoding="utf-8") as fh:
            doc = json.load(fh)
    except (OSError, ValueError) as e:
        raise DocError(f"cannot read {json_path}: {e}") from None
    if not isinstance(doc, dict) or "blocks" not in doc:
        raise DocError(f"{json_path} is not a God-Agent document")
    return doc


def _save(json_path: str, doc: dict) -> None:
    doc["updated"] = time.strftime("%Y-%m-%dT%H:%M:%S")
    tmp = json_path + ".tmp"
    os.makedirs(os.path.dirname(os.path.abspath(json_path)) or ".", exist_ok=True)
    with open(tmp, "w", encoding="utf-8") as fh:
        json.dump(doc, fh, indent=2, ensure_ascii=False)
    os.replace(tmp, json_path)  # atomic: never leave a half-written document


def _cfg(rt) -> dict:
    d = rt.cfg.get("docs") or {}
    root = os.path.expanduser(rt.cfg.get("state", {}).get("root", "~/.god-agent"))
    return {
        "enabled": bool(d.get("enabled", True)),
        "default_format": str(d.get("default_format") or "md"),
        "embed_images": bool(d.get("embed_images", True)),
        "pdf_format": str(d.get("pdf_format") or "A4"),
        "max_image_mb": int(d.get("max_image_mb") or 25),
        "output_dir": os.path.expanduser(d.get("output_dir") or ""),
    }


def _err(rt, action: str, exc: BaseException) -> str:
    msg = str(exc) if isinstance(exc, DocError) else f"{type(exc).__name__}: {exc}"
    try:
        rt.record(action, f"ERROR: {msg}", {"error": type(exc).__name__})
    except Exception:
        pass
    return f"ERROR: {msg}"


def _guard(rt, conf: dict) -> Optional[str]:
    if not conf["enabled"]:
        return "ERROR: document tools are disabled (set docs.enabled=true)"
    return None


def _ingest_image(src: str, assets_dir: str, max_mb: int) -> str:
    """Copy an image into the document's assets dir; return its file name."""
    src = os.path.expanduser(src)
    if not os.path.isfile(src):
        raise DocError(f"no such image: {src}")
    size_mb = os.path.getsize(src) / (1024 * 1024)
    if size_mb > max_mb:
        raise DocError(f"image is {size_mb:.1f}MB, over the {max_mb}MB cap "
                       f"(docs.max_image_mb)")
    os.makedirs(assets_dir, exist_ok=True)
    name = os.path.basename(src)
    dst = os.path.join(assets_dir, name)
    if os.path.abspath(src) != os.path.abspath(dst):
        # Don't silently clobber a different image with the same file name.
        if os.path.exists(dst) and not _same_file(src, dst):
            stem, ext = os.path.splitext(name)
            name = f"{stem}-{int(time.time())}{ext}"
            dst = os.path.join(assets_dir, name)
        shutil.copy2(src, dst)
    return name


def _same_file(a: str, b: str) -> bool:
    try:
        return os.path.samefile(a, b)
    except OSError:
        return False


def _append(json_path: str, block: dict) -> tuple[dict, int]:
    doc = _load(json_path)
    doc.setdefault("blocks", []).append(block)
    _save(json_path, doc)
    return doc, len(doc["blocks"]) - 1


# ---------------------------------------------------------------------------
# Tools: creation
# ---------------------------------------------------------------------------
def doc_create(reg, name: str, args: dict) -> str:
    """Create a new document. Returns the path to use with the other doc tools."""
    rt = get_runtime()
    conf = _cfg(rt)
    if bad := _guard(rt, conf):
        return bad
    path = str(args.get("path", "")).strip()
    if not path:
        return "ERROR: path required"
    title = str(args.get("title", "") or "Untitled document")
    json_path, assets_dir = _paths(path)

    if os.path.isfile(json_path) and not args.get("overwrite"):
        return (f"ERROR: document already exists: {json_path} "
                f"(pass overwrite=true to replace it)")

    doc = {
        "version": SCHEMA_VERSION,
        "title": title,
        "created": time.strftime("%Y-%m-%dT%H:%M:%S"),
        "updated": time.strftime("%Y-%m-%dT%H:%M:%S"),
        "assets": os.path.basename(assets_dir),
        "blocks": [],
    }
    try:
        os.makedirs(assets_dir, exist_ok=True)
        _save(json_path, doc)
    except OSError as e:
        return _err(rt, "doc_create", DocError(f"cannot create {json_path}: {e}"))
    rt.record("doc_create", f"created document {json_path} ('{title}')",
              {"path": json_path, "title": title})
    return f"ok: created {json_path}\nassets: {assets_dir}"


def doc_add_heading(reg, name: str, args: dict) -> str:
    """Add a heading (level 1-6) to a document."""
    rt = get_runtime()
    if bad := _guard(rt, _cfg(rt)):
        return bad
    text = str(args.get("text", "")).strip()
    if not text:
        return "ERROR: text required"
    level = max(1, min(int(args.get("level", 2) or 2), 6))
    try:
        json_path, _ = _paths(str(args.get("path", "")))
        doc, idx = _append(json_path, {"kind": "heading", "level": level, "text": text})
    except Exception as exc:  # noqa: BLE001
        return _err(rt, "doc_add_heading", exc)
    rt.record("doc_add_heading", f"added h{level} to {json_path}",
              {"path": json_path, "level": level})
    return f"ok: block {idx} (heading h{level}) — {len(doc['blocks'])} block(s) total"


def doc_add_text(reg, name: str, args: dict) -> str:
    """Add a paragraph of text to a document."""
    rt = get_runtime()
    if bad := _guard(rt, _cfg(rt)):
        return bad
    text = str(args.get("text", "")).strip()
    if not text:
        return "ERROR: text required"
    try:
        json_path, _ = _paths(str(args.get("path", "")))
        doc, idx = _append(json_path, {"kind": "text", "text": text})
    except Exception as exc:  # noqa: BLE001
        return _err(rt, "doc_add_text", exc)
    rt.record("doc_add_text", f"added paragraph to {json_path}", {"path": json_path})
    return f"ok: block {idx} (text) — {len(doc['blocks'])} block(s) total"


def doc_add_code(reg, name: str, args: dict) -> str:
    """Add a fenced code block to a document."""
    rt = get_runtime()
    if bad := _guard(rt, _cfg(rt)):
        return bad
    code = str(args.get("code", ""))
    if not code.strip():
        return "ERROR: code required"
    language = str(args.get("language", "") or "")
    try:
        json_path, _ = _paths(str(args.get("path", "")))
        doc, idx = _append(json_path, {"kind": "code", "text": code,
                                       "language": language})
    except Exception as exc:  # noqa: BLE001
        return _err(rt, "doc_add_code", exc)
    rt.record("doc_add_code", f"added code block to {json_path}", {"path": json_path})
    return f"ok: block {idx} (code) — {len(doc['blocks'])} block(s) total"


def doc_add_image(reg, name: str, args: dict) -> str:
    """Embed an image in a document — the point of the whole toolset.

    The image is copied into the document's assets directory, so the document
    stays portable. Pair with `browser_screenshot` to capture, then embed, or
    with `image_generate` to draw a diagram and drop it straight in.
    """
    rt = get_runtime()
    conf = _cfg(rt)
    if bad := _guard(rt, conf):
        return bad
    image = str(args.get("image", "")).strip()
    if not image:
        return "ERROR: image required (path to an image file)"
    caption = str(args.get("caption", "") or "")
    alt = str(args.get("alt", "") or caption or "image")
    width = args.get("width")
    try:
        json_path, assets_dir = _paths(str(args.get("path", "")))
        if not os.path.isfile(json_path):
            raise DocError(f"no such document: {json_path} (create it with doc_create)")
        fname = _ingest_image(image, assets_dir, conf["max_image_mb"])
        block: dict[str, Any] = {"kind": "image", "path": fname,
                                 "caption": caption, "alt": alt}
        if width:
            block["width"] = int(width)
        doc, idx = _append(json_path, block)
    except Exception as exc:  # noqa: BLE001
        return _err(rt, "doc_add_image", exc)
    rt.record("doc_add_image", f"embedded {fname} in {json_path}",
              {"path": json_path, "image": fname, "caption": caption})
    return (f"ok: block {idx} (image {fname}) — "
            f"{len(doc['blocks'])} block(s) total")


def doc_add_quote(reg, name: str, args: dict) -> str:
    """Add a block quote to a document."""
    rt = get_runtime()
    if bad := _guard(rt, _cfg(rt)):
        return bad
    text = str(args.get("text", "")).strip()
    if not text:
        return "ERROR: text required"
    try:
        json_path, _ = _paths(str(args.get("path", "")))
        doc, idx = _append(json_path, {"kind": "quote", "text": text})
    except Exception as exc:  # noqa: BLE001
        return _err(rt, "doc_add_quote", exc)
    rt.record("doc_add_quote", f"added quote to {json_path}", {"path": json_path})
    return f"ok: block {idx} (quote) — {len(doc['blocks'])} block(s) total"


def doc_add_page_break(reg, name: str, args: dict) -> str:
    """Insert a page break (affects pdf and docx rendering)."""
    rt = get_runtime()
    if bad := _guard(rt, _cfg(rt)):
        return bad
    try:
        json_path, _ = _paths(str(args.get("path", "")))
        doc, idx = _append(json_path, {"kind": "page_break", "text": ""})
    except Exception as exc:  # noqa: BLE001
        return _err(rt, "doc_add_page_break", exc)
    rt.record("doc_add_page_break", f"added page break to {json_path}",
              {"path": json_path})
    return f"ok: block {idx} (page_break) — {len(doc['blocks'])} block(s) total"


# ---------------------------------------------------------------------------
# Tools: inspection and editing
# ---------------------------------------------------------------------------
def doc_outline(reg, name: str, args: dict) -> str:
    """List a document's blocks with their indices, for targeted editing."""
    rt = get_runtime()
    try:
        json_path, _ = _paths(str(args.get("path", "")))
        doc = _load(json_path)
    except Exception as exc:  # noqa: BLE001
        return _err(rt, "doc_outline", exc)
    blocks = doc.get("blocks", [])
    if not blocks:
        return f"(empty document) {json_path}\ntitle: {doc.get('title','')}"
    lines = [f"{os.path.basename(json_path)} — {len(blocks)} block(s), "
             f"title: {doc.get('title','')}", ""]
    for i, b in enumerate(blocks):
        kind = b.get("kind", "?")
        if kind == "image":
            preview = f"{b.get('path','')} {b.get('caption','')}".strip()
        else:
            preview = (b.get("text", "") or "").replace("\n", " ")
        if len(preview) > 70:
            preview = preview[:70] + "..."
        lines.append(f"[{i}] {kind:<11} {preview}")
    return "\n".join(lines)


def doc_edit(reg, name: str, args: dict) -> str:
    """Replace the content of one block, in place.

    Edit by index from `doc_outline` rather than regenerating the document —
    that keeps the surrounding blocks (and image references) intact.
    """
    rt = get_runtime()
    if bad := _guard(rt, _cfg(rt)):
        return bad
    idx = args.get("block")
    if idx is None:
        return "ERROR: block index required (see doc_outline)"
    try:
        json_path, _ = _paths(str(args.get("path", "")))
        doc = _load(json_path)
        blocks = doc.get("blocks", [])
        i = int(idx)
        if i < 0 or i >= len(blocks):
            raise DocError(f"block {i} out of range (0..{len(blocks)-1})")
        block = blocks[i]
        changed = []
        if args.get("text") is not None:
            block["text"] = str(args["text"])
            changed.append("text")
        if args.get("caption") is not None:
            block["caption"] = str(args["caption"])
            changed.append("caption")
        if args.get("level") is not None:
            block["level"] = max(1, min(int(args["level"]), 6))
            changed.append("level")
        if args.get("width") is not None:
            block["width"] = int(args["width"])
            changed.append("width")
        if not changed:
            return "ERROR: nothing to edit (supply text, caption, level, or width)"
        _save(json_path, doc)
    except Exception as exc:  # noqa: BLE001
        return _err(rt, "doc_edit", exc)
    rt.record("doc_edit", f"edited block {i} of {json_path}",
              {"path": json_path, "block": i, "fields": changed})
    return f"ok: edited block {i} ({', '.join(changed)})"


def doc_remove(reg, name: str, args: dict) -> str:
    """Delete a block from a document."""
    rt = get_runtime()
    if bad := _guard(rt, _cfg(rt)):
        return bad
    idx = args.get("block")
    if idx is None:
        return "ERROR: block index required (see doc_outline)"
    try:
        json_path, _ = _paths(str(args.get("path", "")))
        doc = _load(json_path)
        blocks = doc.get("blocks", [])
        i = int(idx)
        if i < 0 or i >= len(blocks):
            raise DocError(f"block {i} out of range (0..{len(blocks)-1})")
        removed = blocks.pop(i)
        _save(json_path, doc)
    except Exception as exc:  # noqa: BLE001
        return _err(rt, "doc_remove", exc)
    rt.record("doc_remove", f"removed block {i} from {json_path}",
              {"path": json_path, "block": i, "kind": removed.get("kind")})
    return f"ok: removed block {i} ({removed.get('kind')}) — {len(blocks)} block(s) left"


# ---------------------------------------------------------------------------
# Rendering
# ---------------------------------------------------------------------------
def _b64(path: str) -> str:
    with open(path, "rb") as fh:
        return base64.b64encode(fh.read()).decode("ascii")


def _mime(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    return ({".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
             ".gif": "image/gif", ".webp": "image/webp", ".bmp": "image/bmp",
             ".svg": "image/svg+xml"}.get(ext)
            or mimetypes.guess_type(path)[0] or "image/png")


def _img_src(assets_dir: str, fname: str, embed: bool) -> Optional[str]:
    full = os.path.join(assets_dir, fname)
    if embed and os.path.isfile(full):
        return f"data:{_mime(fname)};base64,{_b64(full)}"
    return os.path.join(os.path.basename(assets_dir), fname)


def render_markdown(doc: dict, assets_dir: str, embed: bool = False) -> str:
    out: list[str] = []
    if doc.get("title"):
        out += [f"# {doc['title']}", ""]
    for b in doc.get("blocks", []):
        kind = b.get("kind")
        text = b.get("text", "") or ""
        if kind == "heading":
            out += ["#" * max(1, min(int(b.get("level", 2)), 6)) + " " + text, ""]
        elif kind == "text":
            out += [text, ""]
        elif kind == "quote":
            out += ["> " + line for line in text.splitlines()] + [""]
        elif kind == "code":
            out += ["```" + (b.get("language") or ""), text, "```", ""]
        elif kind == "page_break":
            out += ["---", ""]
        elif kind == "image":
            src = _img_src(assets_dir, b.get("path", ""), embed) or b.get("path", "")
            alt = b.get("alt") or b.get("caption") or "image"
            out += [f"![{alt}]({src})", ""]
            if b.get("caption"):
                out += [f"*{b['caption']}*", ""]
    return "\n".join(out).rstrip() + "\n"


def render_html(doc: dict, assets_dir: str, embed: bool = True) -> str:
    title = _html.escape(doc.get("title", "Document"))
    body: list[str] = []
    for b in doc.get("blocks", []):
        kind = b.get("kind")
        text = _html.escape(b.get("text", "") or "")
        if kind == "heading":
            lvl = max(1, min(int(b.get("level", 2)), 6))
            body.append(f"<h{lvl}>{text}</h{lvl}>")
        elif kind == "text":
            body.append(f"<p>{text.replace(chr(10), '<br>')}</p>")
        elif kind == "quote":
            body.append(f"<blockquote>{text.replace(chr(10), '<br>')}</blockquote>")
        elif kind == "code":
            lang = _html.escape(b.get("language") or "")
            cls = f' class="language-{lang}"' if lang else ""
            body.append(f"<pre><code{cls}>{text}</code></pre>")
        elif kind == "page_break":
            body.append('<div class="page-break"></div>')
        elif kind == "image":
            src = _img_src(assets_dir, b.get("path", ""), embed) or b.get("path", "")
            alt = _html.escape(b.get("alt") or b.get("caption") or "image")
            w = f' style="max-width:{int(b["width"])}px"' if b.get("width") else ""
            body.append(
                f"<figure><img src=\"{src}\" alt=\"{alt}\"{w}>"
                + (f"<figcaption>{_html.escape(b['caption'])}</figcaption>"
                   if b.get("caption") else "")
                + "</figure>")
    return (
        "<!DOCTYPE html>\n<html><head><meta charset=\"utf-8\">"
        f"<title>{title}</title><style>"
        "body{font:16px/1.6 -apple-system,Segoe UI,Roboto,Helvetica,Arial,sans-serif;"
        "max-width:820px;margin:2rem auto;padding:0 1.25rem;color:#1a1a1a}"
        "h1,h2,h3{line-height:1.25;margin-top:1.8rem}"
        "pre{background:#f6f8fa;padding:1rem;border-radius:6px;overflow-x:auto}"
        "code{font-family:ui-monospace,SFMono-Regular,Menlo,monospace;font-size:.9em}"
        "blockquote{border-left:4px solid #d0d7de;margin:1rem 0;padding:.25rem 1rem;color:#57606a}"
        "figure{margin:1.5rem 0}figure img{max-width:100%;height:auto;border:1px solid #d0d7de;"
        "border-radius:6px}figcaption{font-size:.9em;color:#57606a;margin-top:.5rem;text-align:center}"
        ".page-break{page-break-after:always}"
        "</style></head><body>\n"
        f"<h1>{title}</h1>\n" + "\n".join(body) + "\n</body></html>\n"
    )


def render_pdf(html_str: str, out_path: str, page_format: str = "A4") -> None:
    """Render HTML to PDF through the headless Chromium the browser tools use."""
    try:
        from playwright.sync_api import sync_playwright  # noqa: PLC0415
    except ImportError as exc:
        raise DocError(
            "PDF rendering needs Playwright (already required by the browser "
            "tools): pip install 'god-agent[browser]' && playwright install chromium"
        ) from exc
    pw = sync_playwright().start()
    try:
        browser = pw.chromium.launch(
            args=["--no-sandbox", "--disable-dev-shm-usage"])
        try:
            page = browser.new_page()
            page.set_content(html_str, wait_until="load")
            page.pdf(path=out_path, format=page_format, print_background=True)
        finally:
            browser.close()
    finally:
        pw.stop()


def render_docx(doc: dict, assets_dir: str, out_path: str) -> None:
    try:
        import docx  # noqa: PLC0415
    except ImportError as exc:
        raise DocError(
            "DOCX rendering needs python-docx: pip install 'god-agent[docs]'"
        ) from exc
    from docx.shared import Pt  # noqa: PLC0415

    d = docx.Document()
    if doc.get("title"):
        d.add_heading(doc["title"], level=0)
    for b in doc.get("blocks", []):
        kind = b.get("kind")
        text = b.get("text", "") or ""
        if kind == "heading":
            d.add_heading(text, level=max(1, min(int(b.get("level", 2)), 6)))
        elif kind in ("text", "quote"):
            p = d.add_paragraph(text)
            if kind == "quote":
                p.style = "Intense Quote" if "Intense Quote" in [
                    s.name for s in d.styles] else "Quote"
        elif kind == "code":
            p = d.add_paragraph(text)
            p.runs[0].font.name = "Courier New"
            p.runs[0].font.size = Pt(10)
        elif kind == "page_break":
            d.add_page_break()
        elif kind == "image":
            full = os.path.join(assets_dir, b.get("path", ""))
            if os.path.isfile(full):
                d.add_picture(full)
            if b.get("caption"):
                cap = d.add_paragraph(b["caption"])
                cap.alignment = 1
    d.save(out_path)


def doc_render(reg, name: str, args: dict) -> str:
    """Render a document to markdown, HTML, PDF, or DOCX.

    One document renders to any format, so you can re-render after editing
    without rebuilding it. Images are embedded in HTML/PDF output so the
    result is a single portable file.
    """
    rt = get_runtime()
    conf = _cfg(rt)
    if bad := _guard(rt, conf):
        return bad
    fmt = str(args.get("format") or conf["default_format"]).lower()
    if fmt not in RENDER_FORMATS:
        return f"ERROR: format must be one of {'|'.join(RENDER_FORMATS)}"
    try:
        json_path, assets_dir = _paths(str(args.get("path", "")))
        doc = _load(json_path)
    except Exception as exc:  # noqa: BLE001
        return _err(rt, "doc_render", exc)

    out = str(args.get("out", "")).strip()
    if not out:
        base = json_path[: -len(SUFFIX)]
        out = f"{base}.{fmt}"
    out = os.path.expanduser(out)
    try:
        os.makedirs(os.path.dirname(os.path.abspath(out)) or ".", exist_ok=True)
        if fmt == "md":
            # Markdown always references the assets directory relatively:
            # that is the format's convention, it keeps the file readable and
            # diffable, and the assets sit right next to it so it stays
            # portable. Base64 belongs in HTML/PDF, which are single-file
            # outputs that must survive being emailed.
            with open(out, "w", encoding="utf-8") as fh:
                fh.write(render_markdown(doc, assets_dir, embed=False))
        elif fmt == "html":
            with open(out, "w", encoding="utf-8") as fh:
                fh.write(render_html(doc, assets_dir, embed=conf["embed_images"]))
        elif fmt == "pdf":
            render_pdf(render_html(doc, assets_dir, embed=True), out,
                       conf["pdf_format"])
        else:
            render_docx(doc, assets_dir, out)
    except Exception as exc:  # noqa: BLE001
        return _err(rt, "doc_render", exc)

    size = os.path.getsize(out) if os.path.isfile(out) else 0
    rt.record("doc_render", f"rendered {json_path} -> {out} ({fmt}, {size} bytes)",
              {"path": json_path, "out": out, "format": fmt, "bytes": size})
    return f"ok: rendered {fmt} -> {out} ({size} bytes)"


__all__ = [
    "doc_create", "doc_add_heading", "doc_add_text", "doc_add_code",
    "doc_add_image", "doc_add_quote", "doc_add_page_break",
    "doc_outline", "doc_edit", "doc_remove", "doc_render",
    "render_markdown", "render_html", "render_pdf", "render_docx",
    "DocError", "RENDER_FORMATS", "BLOCK_KINDS",
]

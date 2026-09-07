"""Tests for document creation/editing and image editing.

These use the real Pillow and python-docx where available, so a DOCX is
verified by opening it back up rather than by checking that a file exists.
PDF rendering is skipped when the Chromium binary is not installed.
"""
import base64
import json
import os
import sys
import tempfile
from unittest import mock

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "tests"))

try:
    import pytest  # noqa: E402
except ImportError:  # pragma: no cover
    pytest = None

from god_agent.config import default_config  # noqa: E402
from god_agent.policy import TOOL_RISK  # noqa: E402
from god_agent.runtime import Runtime  # noqa: E402
from god_agent.tools import docs as D  # noqa: E402
from god_agent.tools import imaging as I  # noqa: E402

from test_browser import make_cfg, make_rt  # noqa: E402

DOC_TOOLS = ["doc_create", "doc_add_heading", "doc_add_text", "doc_add_code",
             "doc_add_image", "doc_add_quote", "doc_add_page_break",
             "doc_outline", "doc_edit", "doc_remove", "doc_render"]
IMAGE_TOOLS = ["image_info", "image_edit", "image_compose"]


def _png(path, w=120, h=80, color=(200, 30, 30)):
    """Write a real PNG using Pillow."""
    try:
        from PIL import Image
    except ImportError:  # pragma: no cover
        with open(path, "wb") as fh:
            fh.write(base64.b64decode(
                "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8AAAwAB/wD/9p0AAAAASUVORK5CYII="))
        return path
    Image.new("RGB", (w, h), color).save(path)
    return path


def _build_report(rt, d, with_image=True):
    """Create a small document and return its json path."""
    reg = rt.registry
    doc = os.path.join(d, "report")
    assert D.doc_create(reg, "doc_create", {"path": doc, "title": "Incident Report"}).startswith("ok:")
    D.doc_add_heading(reg, "doc_add_heading", {"path": doc, "text": "Summary", "level": 2})
    D.doc_add_text(reg, "doc_add_text", {"path": doc, "text": "Disk filled up on db-01."})
    D.doc_add_code(reg, "doc_add_code", {"path": doc, "code": "df -h", "language": "bash"})
    if with_image:
        shot = _png(os.path.join(d, "shot.png"))
        res = D.doc_add_image(reg, "doc_add_image",
                              {"path": doc, "image": shot, "caption": "Disk usage"})
        assert res.startswith("ok:"), res
    return doc


# ---------------------------------------------------------------------------
# Registration
# ---------------------------------------------------------------------------
def test_all_doc_and_image_tools_registered():
    with tempfile.TemporaryDirectory() as d:
        rt = make_rt(d)
        names = {t.name for t in rt.registry.all()}
        for tool in DOC_TOOLS + IMAGE_TOOLS:
            assert tool in names, f"{tool} not registered"


def test_tools_are_risk_graded():
    for tool in DOC_TOOLS + IMAGE_TOOLS:
        assert tool in TOOL_RISK, f"{tool} is ungraded"
    # Reading is low risk; anything that writes to disk is a 3.
    assert TOOL_RISK["doc_outline"] == 1
    assert TOOL_RISK["image_info"] == 1
    assert TOOL_RISK["doc_edit"] == 3
    assert TOOL_RISK["doc_remove"] == 3
    assert TOOL_RISK["image_edit"] == 3


# ---------------------------------------------------------------------------
# Creation and blocks
# ---------------------------------------------------------------------------
def test_doc_create_makes_document_and_assets():
    with tempfile.TemporaryDirectory() as d:
        rt = make_rt(d)
        doc = os.path.join(d, "r")
        out = D.doc_create(rt.registry, "doc_create", {"path": doc, "title": "T"})
        assert out.startswith("ok:")
        assert os.path.isfile(doc + ".gdoc.json")
        assert os.path.isdir(doc + ".assets")


def test_doc_create_refuses_to_clobber():
    with tempfile.TemporaryDirectory() as d:
        rt = make_rt(d)
        doc = os.path.join(d, "r")
        D.doc_create(rt.registry, "doc_create", {"path": doc, "title": "T"})
        out = D.doc_create(rt.registry, "doc_create", {"path": doc, "title": "T2"})
        assert out.startswith("ERROR:") and "already exists" in out
        # ...but an explicit overwrite is honoured.
        out2 = D.doc_create(rt.registry, "doc_create",
                            {"path": doc, "title": "T2", "overwrite": True})
        assert out2.startswith("ok:")


def test_doc_create_requires_title_and_path():
    with tempfile.TemporaryDirectory() as d:
        rt = make_rt(d)
        assert D.doc_create(rt.registry, "doc_create", {}).startswith("ERROR:")


def test_path_suffix_is_normalised():
    """Passing the .gdoc.json path or the bare name must refer to one document."""
    with tempfile.TemporaryDirectory() as d:
        rt = make_rt(d)
        D.doc_create(rt.registry, "doc_create",
                     {"path": os.path.join(d, "r"), "title": "T"})
        D.doc_add_text(rt.registry, "doc_add_text",
                       {"path": os.path.join(d, "r.gdoc.json"), "text": "hello"})
        out = D.doc_outline(rt.registry, "doc_outline", {"path": os.path.join(d, "r")})
        assert "hello" in out


# ---------------------------------------------------------------------------
# Images
# ---------------------------------------------------------------------------
def test_add_image_copies_into_assets():
    with tempfile.TemporaryDirectory() as d:
        rt = make_rt(d)
        doc = os.path.join(d, "r")
        D.doc_create(rt.registry, "doc_create", {"path": doc, "title": "T"})
        shot = _png(os.path.join(d, "shot.png"))
        D.doc_add_image(rt.registry, "doc_add_image",
                        {"path": doc, "image": shot, "caption": "c"})
        copied = os.path.join(doc + ".assets", "shot.png")
        assert os.path.isfile(copied), "image must travel with the document"
        data = json.load(open(doc + ".gdoc.json"))
        img = [b for b in data["blocks"] if b["kind"] == "image"][0]
        assert img["path"] == "shot.png"   # stored relative, not absolute
        assert img["caption"] == "c"


def test_add_image_rejects_missing_file():
    with tempfile.TemporaryDirectory() as d:
        rt = make_rt(d)
        doc = os.path.join(d, "r")
        D.doc_create(rt.registry, "doc_create", {"path": doc, "title": "T"})
        out = D.doc_add_image(rt.registry, "doc_add_image",
                              {"path": doc, "image": os.path.join(d, "nope.png")})
        assert out.startswith("ERROR:") and "no such image" in out


def test_add_image_to_missing_document():
    with tempfile.TemporaryDirectory() as d:
        rt = make_rt(d)
        shot = _png(os.path.join(d, "s.png"))
        out = D.doc_add_image(rt.registry, "doc_add_image",
                              {"path": os.path.join(d, "ghost"), "image": shot})
        assert out.startswith("ERROR:") and "no such document" in out


def test_same_name_images_do_not_clobber():
    with tempfile.TemporaryDirectory() as d:
        rt = make_rt(d)
        doc = os.path.join(d, "r")
        D.doc_create(rt.registry, "doc_create", {"path": doc, "title": "T"})
        os.makedirs(os.path.join(d, "a"), exist_ok=True)
        os.makedirs(os.path.join(d, "b"), exist_ok=True)
        a = _png(os.path.join(d, "a", "shot.png"), color=(255, 0, 0))
        b = _png(os.path.join(d, "b", "shot.png"), color=(0, 0, 255))
        D.doc_add_image(rt.registry, "doc_add_image", {"path": doc, "image": a})
        out = D.doc_add_image(rt.registry, "doc_add_image", {"path": doc, "image": b})
        assert out.startswith("ok:")
        names = os.listdir(doc + ".assets")
        assert len(names) == 2, f"second image overwrote the first: {names}"


# ---------------------------------------------------------------------------
# Editing
# ---------------------------------------------------------------------------
def test_outline_lists_blocks_with_indices():
    with tempfile.TemporaryDirectory() as d:
        rt = make_rt(d)
        doc = _build_report(rt, d, with_image=False)
        out = D.doc_outline(rt.registry, "doc_outline", {"path": doc})
        assert "[0]" in out and "[1]" in out and "[2]" in out
        assert "Summary" in out
        assert "3 block(s)" in out


def test_outline_on_empty_document():
    with tempfile.TemporaryDirectory() as d:
        rt = make_rt(d)
        doc = os.path.join(d, "r")
        D.doc_create(rt.registry, "doc_create", {"path": doc, "title": "T"})
        assert "empty document" in D.doc_outline(rt.registry, "doc_outline", {"path": doc})


def test_edit_replaces_block_in_place():
    with tempfile.TemporaryDirectory() as d:
        rt = make_rt(d)
        doc = _build_report(rt, d, with_image=False)
        out = D.doc_edit(rt.registry, "doc_edit",
                         {"path": doc, "block": 1, "text": "Memory exhausted on db-02."})
        assert out.startswith("ok:")
        data = json.load(open(doc + ".gdoc.json"))
        assert data["blocks"][1]["text"] == "Memory exhausted on db-02."
        # Neighbours untouched — that's the point of index-based editing.
        assert data["blocks"][0]["text"] == "Summary"


def test_edit_rejects_out_of_range():
    with tempfile.TemporaryDirectory() as d:
        rt = make_rt(d)
        doc = _build_report(rt, d, with_image=False)
        out = D.doc_edit(rt.registry, "doc_edit",
                         {"path": doc, "block": 99, "text": "x"})
        assert out.startswith("ERROR:") and "out of range" in out


def test_edit_with_nothing_to_change():
    with tempfile.TemporaryDirectory() as d:
        rt = make_rt(d)
        doc = _build_report(rt, d, with_image=False)
        out = D.doc_edit(rt.registry, "doc_edit", {"path": doc, "block": 0})
        assert out.startswith("ERROR:") and "nothing to edit" in out


def test_remove_deletes_block():
    with tempfile.TemporaryDirectory() as d:
        rt = make_rt(d)
        doc = _build_report(rt, d, with_image=False)
        out = D.doc_remove(rt.registry, "doc_remove", {"path": doc, "block": 1})
        assert out.startswith("ok:") and "2 block(s) left" in out
        data = json.load(open(doc + ".gdoc.json"))
        assert len(data["blocks"]) == 2


def test_edit_is_reflected_on_re_render():
    """Re-rendering after an edit must show the edit — no rebuild needed."""
    with tempfile.TemporaryDirectory() as d:
        rt = make_rt(d)
        doc = _build_report(rt, d, with_image=False)
        md = os.path.join(d, "out.md")
        D.doc_render(rt.registry, "doc_render", {"path": doc, "format": "md", "out": md})
        assert "Disk filled up" in open(md).read()
        D.doc_edit(rt.registry, "doc_edit",
                   {"path": doc, "block": 1, "text": "Memory exhausted."})
        D.doc_render(rt.registry, "doc_render", {"path": doc, "format": "md", "out": md})
        text = open(md).read()
        assert "Memory exhausted." in text
        assert "Disk filled up" not in text


# ---------------------------------------------------------------------------
# Rendering
# ---------------------------------------------------------------------------
def test_render_markdown_structure():
    with tempfile.TemporaryDirectory() as d:
        rt = make_rt(d)
        doc = _build_report(rt, d)
        md = os.path.join(d, "out.md")
        D.doc_render(rt.registry, "doc_render", {"path": doc, "format": "md", "out": md})
        text = open(md).read()
        assert text.startswith("# Incident Report")
        assert "## Summary" in text
        assert "```bash" in text and "df -h" in text
        assert "![Disk usage](" in text
        # Markdown must stay readable: reference the file, don't inline it.
        assert "base64," not in text
        assert "report.assets/shot.png" in text


def test_render_html_embeds_image_as_data_uri():
    with tempfile.TemporaryDirectory() as d:
        rt = make_rt(d)
        doc = _build_report(rt, d)
        html_out = os.path.join(d, "out.html")
        D.doc_render(rt.registry, "doc_render",
                     {"path": doc, "format": "html", "out": html_out})
        text = open(html_out).read()
        assert "<h1>Incident Report</h1>" in text
        # Embedded, so the file is self-contained and portable.
        assert "data:image/png;base64," in text
        assert "<figcaption>Disk usage</figcaption>" in text


def test_render_html_can_reference_files_instead():
    with tempfile.TemporaryDirectory() as d:
        rt = make_rt(d)
        rt.cfg["docs"]["embed_images"] = False
        doc = _build_report(rt, d)
        html_out = os.path.join(d, "out.html")
        D.doc_render(rt.registry, "doc_render",
                     {"path": doc, "format": "html", "out": html_out})
        text = open(html_out).read()
        assert "data:image/png;base64," not in text
        # _build_report names the document "report", so that is the assets dir.
        assert "report.assets/shot.png" in text


def test_render_docx_produces_real_document():
    try:
        import docx  # noqa: F401
    except ImportError:  # pragma: no cover
        if pytest:
            pytest.skip("python-docx not installed")
        return
    with tempfile.TemporaryDirectory() as d:
        rt = make_rt(d)
        doc = _build_report(rt, d)
        out = os.path.join(d, "out.docx")
        res = D.doc_render(rt.registry, "doc_render",
                           {"path": doc, "format": "docx", "out": out})
        assert res.startswith("ok:"), res
        # Open it back up and confirm the structure really is there.
        import docx as dx

        d2 = dx.Document(out)
        texts = [p.text for p in d2.paragraphs]
        assert any("Incident Report" in t for t in texts)
        assert any("Summary" in t for t in texts)
        assert any("Disk filled up" in t for t in texts)
        assert any("Disk usage" in t for t in texts)
        assert len(d2.inline_shapes) >= 1, "image did not make it into the docx"


def test_render_rejects_bad_format():
    with tempfile.TemporaryDirectory() as d:
        rt = make_rt(d)
        doc = _build_report(rt, d, with_image=False)
        out = D.doc_render(rt.registry, "doc_render", {"path": doc, "format": "rtf"})
        assert out.startswith("ERROR:") and "format must be" in out


def test_render_missing_document():
    with tempfile.TemporaryDirectory() as d:
        rt = make_rt(d)
        out = D.doc_render(rt.registry, "doc_render",
                           {"path": os.path.join(d, "ghost"), "format": "md"})
        assert out.startswith("ERROR:") and "no such document" in out


def test_render_defaults_output_path_beside_document():
    with tempfile.TemporaryDirectory() as d:
        rt = make_rt(d)
        doc = _build_report(rt, d, with_image=False)
        D.doc_render(rt.registry, "doc_render", {"path": doc, "format": "md"})
        assert os.path.isfile(doc + ".md")


def test_render_pdf_if_chromium_present():
    try:
        from playwright.sync_api import sync_playwright  # noqa: PLC0415

        pw = sync_playwright().start()
        try:
            pw.chromium.launch(headless=True)
        finally:
            pw.stop()
    except Exception as exc:  # noqa: BLE001
        if pytest is not None:
            pytest.skip(f"no Chromium binary available: {exc}")
        return

    with tempfile.TemporaryDirectory() as d:
        rt = make_rt(d)
        doc = _build_report(rt, d)
        out = os.path.join(d, "out.pdf")
        res = D.doc_render(rt.registry, "doc_render",
                           {"path": doc, "format": "pdf", "out": out})
        assert res.startswith("ok:"), res
        with open(out, "rb") as fh:
            assert fh.read(5) == b"%PDF-", "not a PDF"


def test_documents_are_audited():
    with tempfile.TemporaryDirectory() as d:
        rt = make_rt(d)
        _build_report(rt, d)
        with open(rt.cfg["audit"]["path"]) as fh:
            blob = fh.read()
        for action in ("doc_create", "doc_add_text", "doc_add_image"):
            assert action in blob


def test_docs_disabled_blocks_writes():
    with tempfile.TemporaryDirectory() as d:
        rt = make_rt(d)
        rt.cfg["docs"]["enabled"] = False
        for fn in (D.doc_create, D.doc_add_text, D.doc_edit, D.doc_remove):
            out = fn(rt.registry, fn.__name__, {"path": "/tmp/x", "text": "t",
                                                "block": 0, "title": "t"})
            assert "disabled" in out, f"{fn.__name__}: {out}"


# ---------------------------------------------------------------------------
# Image editing
# ---------------------------------------------------------------------------
def test_image_info_reports_dimensions():
    try:
        from PIL import Image  # noqa: F401
    except ImportError:  # pragma: no cover
        if pytest:
            pytest.skip("Pillow not installed")
        return
    with tempfile.TemporaryDirectory() as d:
        rt = make_rt(d)
        p = _png(os.path.join(d, "a.png"), w=120, h=80)
        out = I.image_info(rt.registry, "image_info", {"path": p})
        assert "120 x 80" in out
        assert "PNG" in out


def test_image_info_missing_file():
    with tempfile.TemporaryDirectory() as d:
        rt = make_rt(d)
        assert I.image_info(rt.registry, "image_info",
                            {"path": os.path.join(d, "nope.png")}).startswith("ERROR:")


def test_image_edit_resize():
    try:
        from PIL import Image
    except ImportError:  # pragma: no cover
        if pytest:
            pytest.skip("Pillow not installed")
        return
    with tempfile.TemporaryDirectory() as d:
        rt = make_rt(d)
        src = _png(os.path.join(d, "a.png"), w=120, h=80)
        out = os.path.join(d, "b.png")
        res = I.image_edit(rt.registry, "image_edit",
                           {"path": src, "out": out, "width": 60})
        assert res.startswith("ok:"), res
        assert Image.open(out).size == (60, 40)   # aspect ratio preserved
        assert Image.open(src).size == (120, 80)  # source untouched


def test_image_edit_scale_and_grayscale_and_convert():
    try:
        from PIL import Image
    except ImportError:  # pragma: no cover
        if pytest:
            pytest.skip("Pillow not installed")
        return
    with tempfile.TemporaryDirectory() as d:
        rt = make_rt(d)
        src = _png(os.path.join(d, "a.png"), w=100, h=100)
        out = os.path.join(d, "c.jpg")
        res = I.image_edit(rt.registry, "image_edit",
                           {"path": src, "out": out, "scale": 0.5,
                            "grayscale": True, "format": "JPEG"})
        assert res.startswith("ok:"), res
        img = Image.open(out)
        assert img.size == (50, 50)
        assert img.format == "JPEG"
        assert img.mode == "L"


def test_image_edit_crop_and_rotate():
    try:
        from PIL import Image
    except ImportError:  # pragma: no cover
        if pytest:
            pytest.skip("Pillow not installed")
        return
    with tempfile.TemporaryDirectory() as d:
        rt = make_rt(d)
        src = _png(os.path.join(d, "a.png"), w=100, h=100)
        out = os.path.join(d, "d.png")
        res = I.image_edit(rt.registry, "image_edit",
                           {"path": src, "out": out,
                            "crop": "0,0,50,50", "rotate": 90})
        assert res.startswith("ok:"), res
        assert Image.open(out).size == (50, 50)


def test_image_edit_rejects_bad_flip():
    with tempfile.TemporaryDirectory() as d:
        rt = make_rt(d)
        src = _png(os.path.join(d, "a.png"))
        out = I.image_edit(rt.registry, "image_edit",
                           {"path": src, "flip": "sideways"})
        assert out.startswith("ERROR:") and "flip must be" in out


def test_image_edit_defaults_output_name():
    with tempfile.TemporaryDirectory() as d:
        rt = make_rt(d)
        src = _png(os.path.join(d, "a.png"))
        res = I.image_edit(rt.registry, "image_edit", {"path": src, "width": 10})
        assert res.startswith("ok:")
        assert os.path.isfile(os.path.join(d, "a.edited.png"))


def test_image_compose_grid():
    try:
        from PIL import Image
    except ImportError:  # pragma: no cover
        if pytest:
            pytest.skip("Pillow not installed")
        return
    with tempfile.TemporaryDirectory() as d:
        rt = make_rt(d)
        paths = [_png(os.path.join(d, f"{i}.png"), w=40, h=40) for i in range(4)]
        out = os.path.join(d, "grid.png")
        res = I.image_compose(rt.registry, "image_compose",
                              {"paths": paths, "out": out, "layout": "grid",
                               "cols": 2, "spacing": 0})
        assert res.startswith("ok:"), res
        assert Image.open(out).size == (80, 80)


def test_image_compose_horizontal():
    try:
        from PIL import Image
    except ImportError:  # pragma: no cover
        if pytest:
            pytest.skip("Pillow not installed")
        return
    with tempfile.TemporaryDirectory() as d:
        rt = make_rt(d)
        paths = [_png(os.path.join(d, f"{i}.png"), w=30, h=20) for i in range(3)]
        out = os.path.join(d, "row.png")
        I.image_compose(rt.registry, "image_compose",
                        {"paths": paths, "out": out, "layout": "horizontal",
                         "spacing": 0})
        assert Image.open(out).size == (90, 20)


def test_image_compose_auto_columns_and_zero_spacing():
    """Guards a falsy-default bug: cols=0 means 'choose for me', spacing=0
    means 'no gap'. An `or` default silently turns both into something else."""
    try:
        from PIL import Image
    except ImportError:  # pragma: no cover
        if pytest:
            pytest.skip("Pillow not installed")
        return
    with tempfile.TemporaryDirectory() as d:
        rt = make_rt(d)
        paths = [_png(os.path.join(d, f"{i}.png"), w=20, h=20) for i in range(6)]
        out = os.path.join(d, "auto.png")
        res = I.image_compose(rt.registry, "image_compose",
                              {"paths": paths, "out": out, "layout": "grid",
                               "spacing": 0})
        assert res.startswith("ok:"), res
        # 6 images, auto cols -> 3 across, 2 rows, no spacing.
        assert Image.open(out).size == (60, 40)
        assert "3 cols" in res


def test_image_compose_accepts_comma_string():
    with tempfile.TemporaryDirectory() as d:
        rt = make_rt(d)
        a = _png(os.path.join(d, "a.png"))
        b = _png(os.path.join(d, "b.png"))
        res = I.image_compose(rt.registry, "image_compose",
                              {"paths": f"{a},{b}", "out": os.path.join(d, "o.png")})
        assert res.startswith("ok:"), res


def test_image_compose_rejects_empty():
    with tempfile.TemporaryDirectory() as d:
        rt = make_rt(d)
        assert I.image_compose(rt.registry, "image_compose",
                               {"paths": []}).startswith("ERROR:")


def test_missing_pillow_gives_actionable_error():
    with tempfile.TemporaryDirectory() as d:
        rt = make_rt(d)
        p = _png(os.path.join(d, "a.png"))
        with mock.patch.dict(sys.modules, {"PIL": None}):
            out = I.image_info(rt.registry, "image_info", {"path": p})
        assert out.startswith("ERROR:")
        assert "pillow" in out.lower()


def test_imaging_disabled_blocks_tools():
    with tempfile.TemporaryDirectory() as d:
        rt = make_rt(d)
        rt.cfg["imaging"]["enabled"] = False
        p = _png(os.path.join(d, "a.png"))
        for fn, args in ((I.image_info, {"path": p}),
                         (I.image_edit, {"path": p}),
                         (I.image_compose, {"paths": [p]})):
            assert "disabled" in fn(rt.registry, fn.__name__, args)


def test_config_defaults():
    cfg = default_config()
    assert cfg["docs"]["enabled"] is True
    assert cfg["docs"]["default_format"] == "md"
    assert cfg["docs"]["embed_images"] is True
    assert cfg["imaging"]["enabled"] is True
    assert cfg["imaging"]["jpeg_quality"] == 90


if __name__ == "__main__":  # pragma: no cover
    failures = 0
    for name, fn in sorted(globals().items()):
        if name.startswith("test_") and callable(fn):
            try:
                fn()
                print(f"PASS  {name}")
            except Exception as exc:  # noqa: BLE001
                failures += 1
                print(f"FAIL  {name}: {type(exc).__name__}: {exc}")
    print(f"\n{'FAILURES: ' + str(failures) if failures else 'all passed'}")
    sys.exit(1 if failures else 0)
